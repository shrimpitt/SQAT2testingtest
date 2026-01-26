"""
Excel reader utility for reading test data from testdata.xlsx
"""
from openpyxl import load_workbook
from pathlib import Path
from typing import List, Dict, Any


def read_test_data(sheet_name: str = "data") -> List[Dict[str, Any]]:
    """
    Read test data from testdata.xlsx in project root.
    
    Args:
        sheet_name: Name of the sheet to read (default: "data")
        
    Returns:
        List of dictionaries, where each dict represents a row with column headers as keys
    """
    excel_path = Path(__file__).parent.parent / "testdata.xlsx"
    
    if not excel_path.exists():
        raise FileNotFoundError(f"testdata.xlsx not found at {excel_path}")
    
    workbook = load_workbook(excel_path)
    
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found. Available sheets: {workbook.sheetnames}"
        )
    
    worksheet = workbook[sheet_name]
    
    # Extract headers from first row
    headers = []
    for cell in worksheet[1]:
        if cell.value is not None:
            headers.append(cell.value)
    
    if not headers:
        raise ValueError("No headers found in the first row")
    
    # Extract data rows
    test_data = []
    for row in worksheet.iter_rows(min_row=2, values_only=False):
        row_dict = {}
        for i, cell in enumerate(row):
            if i < len(headers):
                # Get the value, handling None values
                value = cell.value
                row_dict[headers[i]] = value if value is not None else ""
        
        # Only add non-empty rows
        if any(row_dict.values()):
            test_data.append(row_dict)
    
    return test_data


def get_test_data_by_name(test_name: str, sheet_name: str = "data") -> Dict[str, Any]:
    """
    Get a specific test data row by test_name.
    
    Args:
        test_name: The test_name value to search for
        sheet_name: Name of the sheet to read (default: "data")
        
    Returns:
        Dictionary representing the matching row, or None if not found
    """
    data = read_test_data(sheet_name)
    for row in data:
        if row.get("test_name") == test_name:
            return row
    return None
